"""Local document readers. No model API calls."""
import re
from pathlib import Path
from typing import List

SUPPORTED_EXTENSIONS = {".pdf", ".docx", ".doc"}


def discover_files(input_dir: Path, recursive: bool) -> List[Path]:
    pattern = "**/*" if recursive else "*"
    files = [
        path
        for path in input_dir.glob(pattern)
        if path.is_file() and path.suffix.lower() in SUPPORTED_EXTENSIONS
    ]
    return sorted(files)


def extract_pdf_text(path: Path, max_pages: int) -> str:
    import pdfplumber

    parts: List[str] = []
    with pdfplumber.open(path) as pdf:
        for page in pdf.pages[:max_pages]:
            text = page.extract_text(x_tolerance=1, y_tolerance=3) or ""
            if text.strip():
                parts.append(text)
    return "\n\n".join(parts)


def extract_docx_text(path: Path) -> str:
    from docx import Document

    document = Document(path)
    parts = [paragraph.text for paragraph in document.paragraphs if paragraph.text.strip()]

    for table in document.tables:
        for row in table.rows:
            row_text = " ".join(cell.text.strip() for cell in row.cells if cell.text.strip())
            if row_text:
                parts.append(row_text)

    return "\n".join(parts)


def extract_doc_text_with_word(path: Path) -> str:
    try:
        import win32com.client  # type: ignore
    except ImportError as exc:
        raise RuntimeError(
            "读取 .doc 文件需要 Microsoft Word 和 pywin32。可先另存为 .docx，或执行：pip install pywin32"
        ) from exc

    word = win32com.client.Dispatch("Word.Application")
    word.Visible = False
    document = None
    try:
        document = word.Documents.Open(str(path.resolve()))
        return document.Content.Text
    finally:
        if document is not None:
            document.Close(False)
        word.Quit()


def extract_text(path: Path, max_pages: int, max_chars: int) -> str:
    suffix = path.suffix.lower()
    if suffix == ".pdf":
        text = extract_pdf_text(path, max_pages)
    elif suffix == ".docx":
        text = extract_docx_text(path)
    elif suffix == ".doc":
        text = extract_doc_text_with_word(path)
    else:
        raise ValueError(f"不支持的文件类型：{suffix}")

    text = normalize_text(text)
    if not text:
        raise RuntimeError("未能从文档中读取到文本，可能是扫描版 PDF 或文档受保护")
    return text[:max_chars]


def normalize_text(text: str) -> str:
    text = text.replace("\r\n", "\n").replace("\r", "\n")
    text = re.sub(r"[ \t]+", " ", text)
    text = re.sub(r"\n{3,}", "\n\n", text)
    return text.strip()


